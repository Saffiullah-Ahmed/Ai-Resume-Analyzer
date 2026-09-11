import sys
import io
import re
from pathlib import Path
import streamlit as st
import docx
from fpdf import FPDF

# Dynamic Paths Setup
CURRENT_DIR = Path(__file__).resolve().parent  # app folder
PROJECT_ROOT = CURRENT_DIR.parent              # project root
SRC_DIR = PROJECT_ROOT / "src"

for path_item in [str(CURRENT_DIR), str(PROJECT_ROOT), str(SRC_DIR)]:
    if path_item not in sys.path:
        sys.path.insert(0, path_item)

# Imports from app folder
from styles import inject_custom_css
from ui_components import (
    render_hero,
    render_feature_card,
    render_metric_card,
    render_score_card,
    render_recommendation_card,
    render_empty_state,
    render_footer,
)

# Backend src imports
from src.pdf_extractor import extract_text_from_pdf
from src.extractor import parse_resume
from src.job_parser import parse_job_description
from src.job_matcher import run_job_matching_pipeline


def clean_filename_string(s: str) -> str:
    """Sanitizes candidate name strings to prevent path errors with raw linebreaks."""
    if not s:
        return "candidate"
    s = re.sub(r'[\r\n\t]+', ' ', str(s)).strip()
    s = re.sub(r'[\\/*?:"<>|]', '', s)
    s = re.sub(r'\s+', '_', s)
    return s.lower()


def extract_text_from_docx(file_stream) -> str:
    """Extracts plain text safely from DOCX binary streams or bytes."""
    try:
        if isinstance(file_stream, bytes):
            buffer = io.BytesIO(file_stream)
        elif hasattr(file_stream, "read"):
            buffer = io.BytesIO(file_stream.read())
        else:
            buffer = file_stream

        doc = docx.Document(buffer)
        return "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    except Exception as e:
        st.error(f"Error reading DOCX file: {e}")
        return ""


def create_docx_report(results: dict) -> bytes:
    """Generates formatted DOCX file bytes for report download."""
    doc = docx.Document()
    doc.add_heading('AI Resume Analyzer - Match Report', level=0)
    
    doc.add_heading('Candidate Overview', level=1)
    doc.add_paragraph(f"Candidate Name: {results.get('candidate_name', 'N/A')}")
    doc.add_paragraph(f"Overall Match Score: {results.get('overall_match_score', 0.0):.1f}%")
    doc.add_paragraph(f"Required Skill Match: {results.get('required_skills_score', 0.0):.1f}%")
    doc.add_paragraph(f"Text Similarity Score: {results.get('similarity_score', 0.0):.1f}%")
    
    doc.add_heading('Missing Skills & Gap Analysis', level=1)
    missing = results.get('missing_skills', [])
    if missing:
        for skill in missing:
            doc.add_paragraph(f"• {skill}")
    else:
        doc.add_paragraph("No critical skill gaps detected.")
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_pdf_report(results: dict) -> bytes:
    """Generates formatted PDF file bytes for report download."""
    pdf = FPDF()
    pdf.add_page()
    
    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 10, "AI Resume Analyzer - Match Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)
    
    pdf.set_font("Helvetica", style="B", size=13)
    pdf.cell(0, 8, "Candidate Overview", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    pdf.cell(0, 6, f"Candidate Name: {results.get('candidate_name', 'N/A')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Overall Match Score: {results.get('overall_match_score', 0.0):.1f}%", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Required Skill Match: {results.get('required_skills_score', 0.0):.1f}%", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Text Similarity Score: {results.get('similarity_score', 0.0):.1f}%", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    pdf.set_font("Helvetica", style="B", size=13)
    pdf.cell(0, 8, "Missing Skills & Gap Analysis", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    
    missing = results.get('missing_skills', [])
    if missing:
        for skill in missing:
            clean_skill = str(skill).encode('latin-1', 'replace').decode('latin-1')
            pdf.cell(0, 6, f"- {clean_skill}", new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.cell(0, 6, "No critical skill gaps detected.", new_x="LMARGIN", new_y="NEXT")
        
    return bytes(pdf.output())


def reset_analysis_state():
    """Resets all session state values."""
    st.session_state["match_results"] = None
    st.session_state["parsed_resume"] = None
    st.session_state["parsed_job"] = None
    st.session_state["temp_pdf_path"] = None
    st.session_state["batch_results"] = None
    st.session_state["resume_key"] = st.session_state.get("resume_key", 0) + 1
    st.session_state["jd_key"] = st.session_state.get("jd_key", 0) + 1


def get_score_card_config(score: float) -> dict:
    """Calculates UI styles depending on overall score threshold."""
    if score >= 75.0:
        return {
            "label": "Strong Match",
            "bg_gradient": "linear-gradient(135deg, #064E3B 0%, #047857 100%)",
            "border_color": "#10B981",
            "text_color": "#FFFFFF",
            "badge_bg": "#059669",
            "badge_text": "#FFFFFF",
            "icon": "⚡"
        }
    elif score >= 50.0:
        return {
            "label": "Moderate Match",
            "bg_gradient": "linear-gradient(135deg, #78350F 0%, #B45309 100%)",
            "border_color": "#F59E0B",
            "text_color": "#FFFFFF",
            "badge_bg": "#D97706",
            "badge_text": "#FFFFFF",
            "icon": "🎯"
        }
    return {
        "label": "Low Match",
        "bg_gradient": "linear-gradient(135deg, #7F1D1D 0%, #B91C1C 100%)",
        "border_color": "#EF4444",
        "text_color": "#FFFFFF",
        "badge_bg": "#DC2626",
        "badge_text": "#FFFFFF",
        "icon": "⚠️"
    }


# Page Config
st.set_page_config(
    page_title="AI Resume Analyzer & Job Matcher",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded"
)

inject_custom_css()

# Session State Initialization
default_session_keys = {
    "match_results": None,
    "parsed_resume": None,
    "parsed_job": None,
    "temp_pdf_path": None,
    "batch_results": None,
    "resume_key": 0,
    "jd_key": 0
}
for key, val in default_session_keys.items():
    if key not in st.session_state:
        st.session_state[key] = val

# -----------------------------------------------------------------------------
# SIDEBAR UI DESIGN
# -----------------------------------------------------------------------------
with st.sidebar:
    st.markdown(
        """
        <div style="display: flex; align-items: center; gap: 12px; margin-bottom: 6px;">
            <div style="
                background: linear-gradient(135deg, #3B82F6 0%, #8B5CF6 100%);
                width: 48px; 
                height: 48px; 
                border-radius: 12px; 
                display: flex; 
                align-items: center; 
                justify-content: center; 
                font-size: 26px;
                box-shadow: 0 4px 14px rgba(59, 130, 246, 0.4);
                flex-shrink: 0;
            ">
                🧠
            </div>
            <div>
                <div style="color: #FFFFFF; font-family: 'Outfit', sans-serif; font-size: 18px; font-weight: 800; line-height: 1.2;">
                    AI Resume Analyzer
                </div>
                <div style="color: #8B5CF6; font-family: 'Outfit', sans-serif; font-size: 18px; font-weight: 800; line-height: 1.2;">
                    & Job Matcher
                </div>
            </div>
        </div>
        <div style="color: #64748B; font-size: 12px; font-weight: 500; margin-top: 6px; margin-left: 60px;">
            Turn Your Skills Into Opportunities
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown('<hr class="sidebar-divider">', unsafe_allow_html=True)

    app_mode = st.radio(
        label="Sidebar Navigation",
        options=[
            "🏠  Dashboard",
            "📄  Resume Analysis",
            "🎯  Job Recommendations"
        ],
        index=0,
        label_visibility="collapsed",
    )

    st.markdown('<hr class="sidebar-divider">', unsafe_allow_html=True)

    st.markdown(
        """
        <div style="display: flex; align-items: center; gap: 8px; margin-bottom: 8px;">
            <span style="color: #38BDF8; font-size: 18px;">ℹ️</span>
            <span style="color: #38BDF8; font-size: 16px; font-weight: 700;">About</span>
        </div>
        <p style="color: #94A3B8; font-size: 13px; line-height: 1.5; margin-bottom: 0px;">
            AI-powered tool to analyze your resume, match with job descriptions and recommend suitable roles.
        </p>
        """,
        unsafe_allow_html=True
    )

    st.markdown('<hr class="sidebar-divider">', unsafe_allow_html=True)

    st.markdown(
        """
        <div style="display: flex; align-items: center; gap: 8px; margin-bottom: 8px;">
            <span style="color: #8B5CF6; font-size: 18px;">⚙️</span>
            <span style="color: #38BDF8; font-size: 16px; font-weight: 700;">Tech Stack</span>
        </div>
        <div class="tech-grid">
            <div class="tech-badge"><span>🐍</span> Python</div>
            <div class="tech-badge"><span>👑</span> Streamlit</div>
            <div class="tech-badge"><span>🤖</span> Scikit-learn</div>
            <div class="tech-badge"><span>🐼</span> Pandas</div>
            <div class="tech-badge tech-badge-full"><span>🧠</span> NLP</div>
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown('<hr class="sidebar-divider">', unsafe_allow_html=True)

    st.markdown(
        """
        <div style="display: flex; align-items: center; gap: 8px; color: #38BDF8; font-size: 13px; font-weight: 600;">
            <span>✨</span> Version 1.0.0
        </div>
        <div style="color: #64748B; font-size: 11px; margin-top: 4px;">
            Better Skills • Better Matches
        </div>
        """,
        unsafe_allow_html=True
    )


# Map Selection Modes
is_dashboard = "Dashboard" in app_mode
is_profile = "Resume Analysis" in app_mode
is_recommendations = "Job Recommendations" in app_mode

render_hero(
    title="🚀 AI Resume Analyzer & Job Matcher",
    subtitle="Upload your resume, analyze your skill profile, and evaluate job fit against targeted positions using NLP algorithmic matching."
)

# -----------------------------------------------------------------------------
# MODE 1: DASHBOARD
# -----------------------------------------------------------------------------
if is_dashboard:
    f_col1, f_col2, f_col3, f_col4 = st.columns([1, 1, 1, 1])
    with f_col1:
        render_feature_card("📄", "Resume Parsing", "Extract contact data and structural skills automatically.")
    with f_col2:
        render_feature_card("🎯", "Job Matching", "Evaluate candidate alignment against job requirements.")
    with f_col3:
        render_feature_card("💡", "Gap Analysis", "Pinpoint required missing competencies directly.")
    with f_col4:
        render_feature_card("⚡", "TF-IDF Engine", "Algorithmic text analysis using scikit-learn models.")

    st.write("")

    col_resume, col_job = st.columns([1, 1], gap="medium")
    
    with col_resume:
        st.markdown('<div class="upload-card-title">📄 Upload Resume</div><p class="upload-card-desc">PDF or DOCX format for automatic parsing.</p>', unsafe_allow_html=True)
        uploaded_resume = st.file_uploader("Choose file", type=["pdf", "docx", "txt"], key=f"resume_uploader_{st.session_state['resume_key']}", label_visibility="collapsed")
        if uploaded_resume:
            temp_dir = PROJECT_ROOT / "data" / "raw"
            temp_dir.mkdir(parents=True, exist_ok=True)
            temp_path = temp_dir / f"temp_{uploaded_resume.name}"
            with open(temp_path, "wb") as f:
                f.write(uploaded_resume.getbuffer())
            st.session_state["temp_pdf_path"] = str(temp_path)
            st.success(f"Loaded: {uploaded_resume.name}")

    with col_job:
        st.markdown('<div class="upload-card-title">💼 Job Description</div><p class="upload-card-desc">Paste text or upload target document file.</p>', unsafe_allow_html=True)
        tab_paste, tab_file = st.tabs(["✍️ Paste Text", "📁 Upload File"])
        pasted_jd, uploaded_jd_file = "", None
        with tab_paste:
            pasted_jd = st.text_area("Job Description Text", placeholder="Paste job requirements here...", height=140, key="pasted_jd_area", label_visibility="collapsed")
        with tab_file:
            uploaded_jd_file = st.file_uploader("Upload JD", type=["txt", "pdf", "docx"], key=f"jd_uploader_{st.session_state['jd_key']}", label_visibility="collapsed")

    st.write("")

    btn_col1, btn_col2 = st.columns([3, 1], gap="small")
    with btn_col1:
        run_analysis = st.button("🔍 Run Algorithmic Matching", type="primary", use_container_width=True)
    with btn_col2:
        if st.button("↻ Reset", type="secondary", use_container_width=True):
            reset_analysis_state()
            st.rerun()

    if run_analysis:
        if not st.session_state.get("temp_pdf_path"):
            st.error("Please upload a resume file first.")
        else:
            with st.status("Running Backend Analysis Pipeline...", expanded=True) as status_container:
                try:
                    resume_path = str(st.session_state["temp_pdf_path"])
                    if resume_path.endswith(".pdf"):
                        resume_text = extract_text_from_pdf(resume_path)
                    elif resume_path.endswith(".docx"):
                        with open(resume_path, "rb") as f:
                            resume_text = extract_text_from_docx(f.read())
                    elif resume_path.endswith(".txt"):
                        with open(resume_path, "r", encoding="utf-8", errors="ignore") as f:
                            resume_text = f.read()
                    else:
                        resume_text = ""

                    jd_text = ""
                    if uploaded_jd_file:
                        if uploaded_jd_file.name.endswith(".pdf"):
                            temp_jd_path = PROJECT_ROOT / "data" / "raw" / f"temp_{uploaded_jd_file.name}"
                            with open(temp_jd_path, "wb") as f: 
                                f.write(uploaded_jd_file.getbuffer())
                            jd_text = extract_text_from_pdf(str(temp_jd_path))
                        elif uploaded_jd_file.name.endswith(".docx"):
                            jd_text = extract_text_from_docx(uploaded_jd_file.getvalue())
                        else:
                            jd_text = uploaded_jd_file.read().decode("utf-8")
                    elif pasted_jd.strip():
                        jd_text = pasted_jd.strip()

                    if not jd_text.strip():
                        status_container.update(label="Analysis failed: No job description text provided.", state="error")
                        st.error("Please provide a target Job Description.")
                        st.stop()

                    parsed_resume = parse_resume(resume_text)
                    parsed_job = parse_job_description(jd_text)

                    # KEY FIX: Attach full raw text inputs to parsed dictionaries so matcher has complete data
                    parsed_resume["raw_text"] = resume_text
                    parsed_job["raw_text"] = jd_text

                    raw_name = parsed_resume.get("name") or parsed_resume.get("candidate_name") or "Candidate"
                    clean_name = clean_filename_string(raw_name)
                    parsed_resume["name"] = clean_name
                    parsed_resume["candidate_name"] = clean_name

                    match_results = run_job_matching_pipeline(resume_input=parsed_resume, job_input=parsed_job)

                    st.session_state["match_results"] = match_results
                    st.session_state["parsed_resume"] = parsed_resume
                    st.session_state["parsed_job"] = parsed_job
                    
                    status_container.update(label="Pipeline processing complete.", state="complete", expanded=False)
                    st.rerun()

                except Exception as e:
                    status_container.update(label="An error occurred during pipeline execution.", state="error")
                    st.error(f"Backend Exception: {e}")

    st.divider()

    res = st.session_state.get("match_results")
    if res and isinstance(res, dict) and "error" not in res:
        dash_col1, dash_col2 = st.columns([1, 2], gap="medium")
        with dash_col1:
            raw_score = float(res.get("overall_match_score", 0.0))
            score_config = get_score_card_config(raw_score)
            
            render_score_card(raw_score, score_config)

            st.markdown("<div style='margin-top: 16px;'></div>", unsafe_allow_html=True)
            file_base = clean_filename_string(res.get('candidate_name', 'candidate'))

            st.download_button(
                label="📄 Download Report (.DOCX)",
                data=create_docx_report(res),
                file_name=f"match_report_{file_base}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

            st.markdown("<div style='margin-top: 6px;'></div>", unsafe_allow_html=True)

            st.download_button(
                label="📕 Download Report (.PDF)",
                data=create_pdf_report(res),
                file_name=f"match_report_{file_base}.pdf",
                mime="application/pdf",
                use_container_width=True
            )

        with dash_col2:
            st.markdown("<div style='font-size: 16px; font-weight: 700; color: #FFFFFF; margin-bottom: 12px;'>Computed Analysis Metrics</div>", unsafe_allow_html=True)
            m_col1, m_col2 = st.columns([1, 1])
            with m_col1:
                render_metric_card("Required Skill Match", f"{res.get('required_skills_score', 0.0):.1f}%", "🎯")
            with m_col2:
                render_metric_card("Text Similarity Score", f"{res.get('similarity_score', 0.0):.1f}%", "📄")

            st.write("")
            render_recommendation_card(res.get("missing_skills", []))
    else:
        render_empty_state()

# -----------------------------------------------------------------------------
# MODE 2: RESUME ANALYSIS
# -----------------------------------------------------------------------------
elif is_profile:
    parsed_prof = st.session_state.get("parsed_resume")

    if parsed_prof and isinstance(parsed_prof, dict):
        cand_name = str(parsed_prof.get("name") or parsed_prof.get("candidate_name") or "Extracted Profile").replace("\n", " ")
        cand_email = parsed_prof.get("email") or "Not Found"
        cand_phone = parsed_prof.get("phone") or "Not Found"
        extracted_skills = parsed_prof.get("skills", [])

        st.markdown(
            f"""
            <div style="background: #1E293B; border: 1px solid #334155; border-radius: 12px; padding: 24px; margin-bottom: 24px;">
                <div style="font-family: 'Outfit', sans-serif; font-size: 22px; font-weight: 800; color: #FFFFFF;">{cand_name}</div>
                <div style="font-size: 14px; color: #94A3B8; margin-top: 6px;">✉️ Email: <strong>{cand_email}</strong> &nbsp;|&nbsp; 📞 Phone: <strong>{cand_phone}</strong></div>
            </div>
            """,
            unsafe_allow_html=True
        )

        st.markdown("<div style='font-size: 18px; font-weight: 700; color: #FFFFFF; margin-bottom: 12px;'>Extracted Skill Set</div>", unsafe_allow_html=True)
        if extracted_skills:
            badges_html = "".join([f'<span style="display: inline-block; background-color: #334155; color: #F8FAFC; padding: 6px 12px; border-radius: 6px; font-size: 13px; font-weight: 600; margin: 0 6px 8px 0; border: 1px solid #475569;">{s}</span>' for s in extracted_skills])
            st.markdown(f"<div>{badges_html}</div>", unsafe_allow_html=True)
        else:
            st.info("No distinct skills extracted from the provided document.")
    else:
        st.warning("No parsed profile data available in memory. Process a resume in 'Dashboard' mode first.")

# -----------------------------------------------------------------------------
# MODE 3: JOB RECOMMENDATIONS
# -----------------------------------------------------------------------------
elif is_recommendations:
    st.markdown("<div style='font-size: 18px; font-weight: 700; color: #FFFFFF; margin-bottom: 8px;'>💼 Job Recommendations & Batch Matching</div>", unsafe_allow_html=True)
    st.info("Compare your active resume against multiple job descriptions at once.")

    if not st.session_state.get("temp_pdf_path"):
        st.warning("Please upload a resume in the 'Dashboard' tab first before running Recommendations.")
    else:
        uploaded_batch_files = st.file_uploader(
            "Upload Multiple Job Descriptions (PDF, DOCX, TXT)", 
            type=["pdf", "docx", "txt"], 
            accept_multiple_files=True,
            key="batch_files_uploader"
        )

        if st.button("🚀 Run Batch Match", type="primary"):
            if not uploaded_batch_files:
                st.error("Please upload at least one job description file.")
            else:
                resume_path = str(st.session_state["temp_pdf_path"])
                if resume_path.endswith(".pdf"):
                    resume_text = extract_text_from_pdf(resume_path)
                elif resume_path.endswith(".docx"):
                    with open(resume_path, "rb") as f:
                        resume_text = extract_text_from_docx(f.read())
                elif resume_path.endswith(".txt"):
                    with open(resume_path, "r", encoding="utf-8", errors="ignore") as f:
                        resume_text = f.read()
                else:
                    resume_text = ""
                
                parsed_resume = parse_resume(resume_text)
                parsed_resume["raw_text"] = resume_text

                results_list = []
                progress_bar = st.progress(0)
                
                for idx, jd_file in enumerate(uploaded_batch_files):
                    if jd_file.name.endswith(".pdf"):
                        temp_jd = PROJECT_ROOT / "data" / "raw" / f"batch_{jd_file.name}"
                        with open(temp_jd, "wb") as f: 
                            f.write(jd_file.getbuffer())
                        jd_text = extract_text_from_pdf(str(temp_jd))
                    elif jd_file.name.endswith(".docx"):
                        jd_text = extract_text_from_docx(jd_file.getvalue())
                    else:
                        jd_text = jd_file.read().decode("utf-8")

                    parsed_job = parse_job_description(jd_text)
                    parsed_job["raw_text"] = jd_text

                    res = run_job_matching_pipeline(resume_input=parsed_resume, job_input=parsed_job)
                    
                    results_list.append({
                        "Job Title / File": jd_file.name,
                        "Overall Score": f"{res.get('overall_match_score', 0.0):.1f}%",
                        "Skill Match": f"{res.get('required_skills_score', 0.0):.1f}%",
                        "Text Similarity": f"{res.get('similarity_score', 0.0):.1f}%",
                        "Missing Skills": ", ".join(res.get("missing_skills", [])) or "None"
                    })
                    progress_bar.progress((idx + 1) / len(uploaded_batch_files))

                st.session_state["batch_results"] = results_list
                st.success("Batch matching completed!")

        if st.session_state.get("batch_results"):
            st.markdown("### Batch Comparison Results")
            st.dataframe(st.session_state["batch_results"], use_container_width=True)

# Global Footer
render_footer(title="AI Resume Analyzer & Job Matcher", tagline="Data-Driven Career Alignment Tool")